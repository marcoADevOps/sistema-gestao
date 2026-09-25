import os

os.environ["DATABASE_URL"] = "sqlite:///./test_reports.db"

from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO

from app.main import app
from app.database import Base, engine, SessionLocal
from app.auth import hash_password
from app import models

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def setup_and_teardown_db():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    db.add(
        models.User(username="admin1", hashed_password=hash_password("adminpass123"), role="admin")
    )
    db.commit()
    db.close()
    yield
    Base.metadata.drop_all(bind=engine)


@pytest.fixture
def client():
    c = TestClient(app)
    c.post("/login", data={"username": "admin1", "password": "adminpass123"})
    return c


def test_export_requires_login():
    anon = TestClient(app)
    response = anon.get("/web/reports/products.xlsx", follow_redirects=False)
    assert response.status_code == 307
    assert response.headers["location"] == "/login"


def test_export_products_excel_contains_product_rows(client):
    client.post("/api/products/", json={"name": "Caderno", "price": 12.5, "stock_quantity": 30})

    response = client.get("/web/reports/products.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    values = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert "Caderno" in values


def test_export_products_pdf_returns_pdf_bytes(client):
    client.post("/api/products/", json={"name": "Lápis", "price": 2.0, "stock_quantity": 5})

    response = client.get("/web/reports/products.pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_export_clients_excel_contains_client_rows(client):
    client.post("/api/clients/", json={"name": "João Silva", "phone": "61999998888"})

    response = client.get("/web/reports/clients.xlsx")
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    values = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert "João Silva" in values


def test_export_clients_pdf_returns_pdf_bytes(client):
    response = client.get("/web/reports/clients.pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_export_sales_excel_contains_sale_item_rows(client):
    product = client.post(
        "/api/products/", json={"name": "Mochila", "price": 100.0, "stock_quantity": 10}
    ).json()
    client.post("/api/sales/", json={"items": [{"product_id": product["id"], "quantity": 3}]})

    response = client.get("/web/reports/sales.xlsx")
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][3] == "Mochila"  # coluna Produto
    assert rows[0][4] == 3  # coluna Quantidade
    assert rows[0][6] == 300.0  # coluna Subtotal


def test_export_sales_pdf_returns_pdf_bytes(client):
    response = client.get("/web/reports/sales.pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_export_sales_excel_filters_by_date_range(client):
    db = SessionLocal()
    product = models.Product(name="Caneta", price=1.0, stock_quantity=100)
    db.add(product)
    db.commit()
    db.refresh(product)

    old_sale = models.Sale(total=1.0, created_at=datetime.utcnow() - timedelta(days=10))
    recent_sale = models.Sale(total=1.0, created_at=datetime.utcnow())
    db.add_all([old_sale, recent_sale])
    db.commit()
    db.refresh(old_sale)
    db.refresh(recent_sale)
    db.add(models.SaleItem(sale_id=old_sale.id, product_id=product.id, quantity=1, unit_price=1.0))
    db.add(models.SaleItem(sale_id=recent_sale.id, product_id=product.id, quantity=1, unit_price=1.0))
    db.commit()
    db.close()

    start_date = (datetime.utcnow() - timedelta(days=1)).date().isoformat()
    response = client.get(f"/web/reports/sales.xlsx?start_date={start_date}")
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1  # só a venda recente entra no filtro


def test_export_sales_with_blank_date_fields_is_not_rejected(client):
    """O formulário do dashboard envia start_date=&end_date= quando os
    campos ficam em branco — isso deve ser tratado como 'sem filtro', não
    como erro de validação."""
    response = client.get("/web/reports/sales.xlsx?start_date=&end_date=")
    assert response.status_code == 200

    response = client.get("/web/reports/sales.pdf?start_date=&end_date=")
    assert response.status_code == 200


def test_export_sales_with_invalid_date_returns_400(client):
    response = client.get("/web/reports/sales.xlsx?start_date=not-a-date")
    assert response.status_code == 400


def test_export_generates_audit_log_entry(client):
    client.get("/web/reports/products.xlsx")

    response = client.get("/web/audit-log")
    assert "export_report" in response.text or "exportado" in response.text
