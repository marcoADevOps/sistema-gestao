"""Geração dos relatórios exportáveis do dashboard (Excel e PDF)."""

from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

from app import models

PAGE_MARGIN = 10


def _autosize_columns(ws) -> None:
    for column_cells in ws.columns:
        length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 40)


def _excel_from_rows(sheet_title: str, headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    _autosize_columns(ws)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _truncate(value, max_len: int) -> str:
    text = str(value)
    return text if len(text) <= max_len else text[: max_len - 1] + "…"


def _pdf_from_rows(
    title: str, headers: list[str], rows: list[list], col_widths: list[int], landscape: bool = False
) -> bytes:
    pdf = FPDF(orientation="L" if landscape else "P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=PAGE_MARGIN)
    pdf.set_margins(PAGE_MARGIN, PAGE_MARGIN, PAGE_MARGIN)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, title, border=0)
    pdf.ln(14)

    pdf.set_font("Helvetica", "B", 9)
    for header, width in zip(headers, col_widths):
        pdf.cell(width, 8, header, border=1)
    pdf.ln(8)

    pdf.set_font("Helvetica", "", 9)
    for row in rows:
        for value, width in zip(row, col_widths):
            pdf.cell(width, 7, str(value), border=1)
        pdf.ln(7)

    if not rows:
        pdf.cell(0, 7, "Nenhum registro encontrado.", border=0)

    return bytes(pdf.output())


def products_to_excel(products: list[models.Product]) -> bytes:
    rows = [
        [p.name, round(p.price, 2), p.stock_quantity, p.created_at.strftime("%d/%m/%Y %H:%M")]
        for p in products
    ]
    return _excel_from_rows("Produtos", ["Nome", "Preço (R$)", "Estoque", "Cadastrado em"], rows)


def products_to_pdf(products: list[models.Product]) -> bytes:
    rows = [
        [_truncate(p.name, 45), f"R$ {p.price:.2f}", p.stock_quantity, p.created_at.strftime("%d/%m/%Y")]
        for p in products
    ]
    return _pdf_from_rows(
        "Relatório de Produtos", ["Nome", "Preço", "Estoque", "Cadastrado em"], rows, [90, 35, 30, 35]
    )


def clients_to_excel(clients: list[models.Client]) -> bytes:
    rows = [
        [c.name, c.phone or "—", c.email or "—", c.created_at.strftime("%d/%m/%Y %H:%M")]
        for c in clients
    ]
    return _excel_from_rows("Clientes", ["Nome", "Telefone", "E-mail", "Cadastrado em"], rows)


def clients_to_pdf(clients: list[models.Client]) -> bytes:
    rows = [
        [_truncate(c.name, 28), c.phone or "-", _truncate(c.email or "-", 32), c.created_at.strftime("%d/%m/%Y")]
        for c in clients
    ]
    return _pdf_from_rows(
        "Relatório de Clientes", ["Nome", "Telefone", "E-mail", "Cadastrado em"], rows, [55, 40, 60, 35]
    )


def _sale_rows(sales: list[models.Sale]) -> list[list]:
    rows = []
    for sale in sales:
        client_name = sale.client.name if sale.client else "Não informado"
        for item in sale.items:
            rows.append(
                [
                    sale.id,
                    sale.created_at.strftime("%d/%m/%Y %H:%M"),
                    client_name,
                    item.product.name,
                    item.quantity,
                    item.unit_price,
                    round(item.quantity * item.unit_price, 2),
                ]
            )
    return rows


def sales_to_excel(sales: list[models.Sale]) -> bytes:
    rows = _sale_rows(sales)
    return _excel_from_rows(
        "Vendas",
        ["Venda", "Data", "Cliente", "Produto", "Quantidade", "Preço unitário (R$)", "Subtotal (R$)"],
        rows,
    )


def sales_to_pdf(sales: list[models.Sale], period_label: str = "") -> bytes:
    rows = [
        [f"#{r[0]}", r[1], _truncate(r[2], 28), _truncate(r[3], 35), r[4], f"R$ {r[5]:.2f}", f"R$ {r[6]:.2f}"]
        for r in _sale_rows(sales)
    ]
    title = "Relatório de Vendas"
    if period_label:
        title += f" — {period_label}"
    return _pdf_from_rows(
        title,
        ["Venda", "Data", "Cliente", "Produto", "Qtd", "Preço unit.", "Subtotal"],
        rows,
        [15, 28, 45, 55, 15, 25, 25],
        landscape=True,
    )
